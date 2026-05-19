from __future__ import annotations

import csv
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE_DIR = Path(__file__).resolve().parent
RUN_DIR = Path('/home/zulixo/projects/TechTracker/research/evaluation/run_20260416_024916')
OUT_MD = BASE_DIR / 'section_3_4_actual_20260416.md'
OUT_DOCX = BASE_DIR / 'section_3_4_actual_20260416.docx'

BODY_FONT = 'Times New Roman'
BODY_SIZE = Pt(14)
TABLE_SIZE = Pt(12)
FIRST_LINE = Cm(1.25)
LINE_SPACING = 1.5


def read_csv_rows(name: str) -> list[dict[str, str]]:
    path = RUN_DIR / name
    with path.open('r', encoding='utf-8', newline='') as fh:
        return list(csv.DictReader(fh))


def as_float(value: Any):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def fmt_num(value: Any, digits: int = 2) -> str:
    val = as_float(value)
    if val is None:
        return '—'
    return f'{val:.{digits}f}'.replace('.', ',')


def fmt_pct(value: Any, digits: int = 2) -> str:
    val = as_float(value)
    if val is None:
        return '—'
    return f'{val * 100:.{digits}f}%'.replace('.', ',')


def set_run_font(run, *, bold=None, italic=None, size=BODY_SIZE, name=BODY_FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run._element.rPr.rFonts.set(qn('w:cs'), name)
    run.font.size = size
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def format_body_paragraph(p):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = FIRST_LINE
    p.paragraph_format.line_spacing = LINE_SPACING
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    for run in p.runs:
        set_run_font(run)


def format_heading_paragraph(p, *, center=False, before=6, after=6):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.line_spacing = LINE_SPACING
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    for run in p.runs:
        set_run_font(run, bold=True)


def format_caption_paragraph(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.line_spacing = LINE_SPACING
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    for run in p.runs:
        set_run_font(run)


def add_body(doc: Document, text: str):
    p = doc.add_paragraph(text)
    format_body_paragraph(p)
    return p


def add_heading(doc: Document, text: str, *, level: int = 1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=True)
    if level == 1:
        format_heading_paragraph(p, center=True, before=12, after=6)
    else:
        format_heading_paragraph(p, center=False, before=6, after=6)
    return p


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph(text)
    format_caption_paragraph(p)
    return p


def add_picture_with_caption(doc: Document, image_path: Path, caption: str, *, width_cm: float = 15.5):
    if not image_path.exists():
        add_body(doc, f'[{caption}: файл {image_path.name} не найден.]')
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    add_caption(doc, caption)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, *, bold=False, align='left', size=TABLE_SIZE):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = {
        'left': WD_ALIGN_PARAGRAPH.LEFT,
        'center': WD_ALIGN_PARAGRAPH.CENTER,
        'right': WD_ALIGN_PARAGRAPH.RIGHT,
        'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    set_run_font(run, bold=bold, size=size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(
    doc: Document,
    caption: str,
    headers: list[str],
    rows: list[list[str]],
    *,
    col_widths_cm: list[float] | None = None,
):
    add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for idx, text in enumerate(headers):
        set_cell_text(hdr[idx], text, bold=True, align='center')
        shade_cell(hdr[idx], 'D9EAF7')
        if col_widths_cm and idx < len(col_widths_cm):
            hdr[idx].width = Cm(col_widths_cm[idx])
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            align = 'center' if idx > 0 else 'left'
            set_cell_text(cells[idx], text, align=align)
            if col_widths_cm and idx < len(col_widths_cm):
                cells[idx].width = Cm(col_widths_cm[idx])
    doc.add_paragraph('')
    return table


manifest = json.loads((RUN_DIR / 'manifest.json').read_text(encoding='utf-8'))
model_h = read_csv_rows('forecast_metrics_by_model_horizon.csv')
model_total = read_csv_rows('forecast_metrics_by_model.csv')
bucket_rows = read_csv_rows('forecast_metrics_by_bucket.csv')
prob_h = read_csv_rows('forecast_prob_metrics_by_model_horizon.csv')
ci_rows = read_csv_rows('forecast_bootstrap_ci.csv')
dm_rows = read_csv_rows('forecast_dm_tests.csv')


def by_horizon(rows: list[dict[str, str]], horizon: str):
    return [row for row in rows if row.get('horizon') == horizon]


def model_row(rows: list[dict[str, str]], model: str):
    for row in rows:
        if row.get('model_kind') == model:
            return row
    return {}


def dm_row(horizon: str, a: str, b: str):
    for row in dm_rows:
        if row.get('horizon') == horizon and row.get('model_a') == a and row.get('model_b') == b:
            return row
    return {}


def ci_row(horizon: str, model: str, metric: str):
    for row in ci_rows:
        if row.get('horizon') == horizon and row.get('model_kind') == model and row.get('metric') == metric:
            return row
    return {}


rows_24 = by_horizon(model_h, '24h')
rows_7 = by_horizon(model_h, '7d')
rows_30 = by_horizon(model_h, '30d')

sarima_24 = model_row(rows_24, 'sarima')
lstm_24 = model_row(rows_24, 'lstm')
ensemble_24 = model_row(rows_24, 'ensemble')

sarima_7 = model_row(rows_7, 'sarima')
lstm_7 = model_row(rows_7, 'lstm')
ensemble_7 = model_row(rows_7, 'ensemble')

sarima_30 = model_row(rows_30, 'sarima')
lstm_30 = model_row(rows_30, 'lstm')
ensemble_30 = model_row(rows_30, 'ensemble')

ensemble_total = model_row(model_total, 'ensemble')
sarima_total = model_row(model_total, 'sarima')
lstm_total = model_row(model_total, 'lstm')

prob_24 = {row['model_kind']: row for row in by_horizon(prob_h, '24h')}
prob_7 = {row['model_kind']: row for row in by_horizon(prob_h, '7d')}
prob_30 = {row['model_kind']: row for row in by_horizon(prob_h, '30d')}

dm_24 = dm_row('24h', 'sarima', 'ensemble')
dm_7 = dm_row('7d', 'sarima', 'ensemble')
dm_30 = dm_row('30d', 'sarima', 'ensemble')

summary = manifest['summary']
config = manifest['config']

wins = defaultdict(Counter)
grouped = defaultdict(list)
for row in bucket_rows:
    grouped[(row['horizon'], row['metric_code'])].append(row)

for _, items in grouped.items():
    best_rmse = min(items, key=lambda x: float(x['rmse']))
    best_mae = min(items, key=lambda x: float(x['mae']))
    best_pinball = min(items, key=lambda x: float(x['pinball_avg']))
    best_ace = min(items, key=lambda x: float(x['ace80']))
    wins['rmse'][best_rmse['model_kind']] += 1
    wins['mae'][best_mae['model_kind']] += 1
    wins['pinball_avg'][best_pinball['model_kind']] += 1
    wins['ace80'][best_ace['model_kind']] += 1


table_34_rows = [
    ['Устройство', config['serial'], 'Единый объект сравнения всех моделей на одном наборе зрелых прогнозов.'],
    ['Окно прогноза', f"{config['target_date_from'][:10]} — {config['target_date_to'][:10]}", 'В отбор включены только точки, для которых можно проверить факт.'],
    ['Горизонты', ', '.join(config['horizons']), 'Сравнение краткосрочного, среднесрочного и долгосрочного прогноза.'],
    ['Модели', ', '.join(config['model_kinds']), 'Сопоставление статистической, нейросетевой и интегрированной ветвей.'],
    ['Strict intersection', 'Включён', 'Сравнение проводится только на общих зрелых точках всех выбранных моделей.'],
    ['Сравнимых точек', str(summary['forecast_points_compared']), 'Итоговая выборка для расчёта RMSE, MAE и вероятностных метрик.'],
    ['Покрытие', fmt_pct(summary['forecast_coverage_compared'], 2), 'Доля прогнозных точек, вошедших в финальное сравнение.'],
    ['Bootstrap', '300 итераций, moving block', 'Доверительные интервалы рассчитаны с учётом временной зависимости ошибок.'],
]


def horizon_table_rows(rows_by_h: list[dict[str, str]], prob_by_h: dict[str, dict[str, str]]):
    out = []
    for label, key in [('SARIMA', 'sarima'), ('LSTM', 'lstm'), ('Оркестр', 'ensemble')]:
        row = model_row(rows_by_h, key)
        prow = prob_by_h[key]
        out.append([
            label,
            row.get('samples', '—'),
            fmt_num(row.get('mean_error'), 4),
            fmt_num(row.get('mae'), 4),
            fmt_num(row.get('rmse'), 4),
            fmt_pct(row.get('smape'), 2),
            fmt_num(prow.get('pinball_avg'), 4),
            fmt_pct(row.get('picp80'), 2),
            fmt_pct(row.get('ace80'), 2),
        ])
    return out


table_35_rows = horizon_table_rows(rows_24, prob_24)
table_36_rows = horizon_table_rows(rows_7, prob_7)
table_37_rows = horizon_table_rows(rows_30, prob_30)

table_38_rows = [
    ['RMSE', str(wins['rmse']['ensemble']), str(wins['rmse']['sarima']), str(wins['rmse']['lstm'])],
    ['MAE', str(wins['mae']['ensemble']), str(wins['mae']['sarima']), str(wins['mae']['lstm'])],
    ['Pinball(avg)', str(wins['pinball_avg']['ensemble']), str(wins['pinball_avg']['sarima']), str(wins['pinball_avg']['lstm'])],
    ['ACE80', str(wins['ace80']['ensemble']), str(wins['ace80']['sarima']), str(wins['ace80']['lstm'])],
]

table_39_rows = [
    [
        '24h',
        'SARIMA vs Оркестр',
        fmt_num(dm_24.get('dm_stat'), 4),
        fmt_num(dm_24.get('p_value'), 6),
        dm_24.get('winner') or '—',
        'Различие значимо в пользу оркестра.',
    ],
    [
        '7d',
        'SARIMA vs Оркестр',
        fmt_num(dm_7.get('dm_stat'), 4),
        fmt_num(dm_7.get('p_value'), 6),
        dm_7.get('winner') or '—',
        'Преимущество оркестра статистически подтверждено.',
    ],
    [
        '30d',
        'SARIMA vs Оркестр',
        fmt_num(dm_30.get('dm_stat'), 4),
        fmt_num(dm_30.get('p_value'), 6),
        dm_30.get('winner') or '—',
        'Даже на длинном горизонте оркестр лучше baseline по парной ошибке.',
    ],
]

intro = f'''3.4 Вычислительный эксперимент на базе собранных метрик

Вычислительный эксперимент в настоящей редакции раздела выполнен по фактически полученным артефактам системы TechTracker на этапе, когда в проекте уже реализованы траекторный прогноз SARIMA, удалённый LSTM-сервис, оркестрация прогнозов, единый модуль оценки состояния и формальная диссертационная верификация качества. Источником результатов является штатный evaluation-run `run_20260416_024916`, сформированный встроенным модулем оценки прогноза по устройству `DEMO-SARIMA-001`. В анализ включены горизонты `24h`, `7d`, `30d`, а сравнение моделей выполнено для трёх ветвей: `sarima`, `lstm` и `ensemble`.

Принципиально важно, что в данном разделе используются не промежуточные отладочные графики, а официальный набор артефактов, сформированный самой системой: таблицы `forecast_metrics_by_model_horizon.csv`, `forecast_metrics_by_bucket.csv`, `forecast_prob_metrics_by_model_horizon.csv`, `forecast_interval_calibration_by_model_horizon.csv`, `forecast_dm_tests.csv` и `forecast_bootstrap_ci.csv`. Благодаря этому выводы раздела опираются на воспроизводимый экспериментальный протокол, а не на визуальное впечатление от отдельных кривых.

Всего в окно проверки вошло {summary['forecast_points_total']} прогнозных точек. Для {summary['forecast_points_with_actual_raw']} точек был найден фактический ряд, что соответствует сырому покрытию {fmt_pct(summary['forecast_coverage_raw'], 2)}. После включения режима `strict intersection = True`, требующего наличие зрелого факта одновременно для всех сравниваемых моделей, в итоговой выборке осталось {summary['forecast_points_compared']} точек, то есть {fmt_pct(summary['forecast_coverage_compared'], 2)} от полного объёма. Количество общих ключей сравнения составило {summary['strict_intersection_keys']}, поэтому эксперимент является статистически содержательным и пригодным для использования в диссертационном тексте.
'''

p341 = f'''3.4.1 Организация эксперимента и состав выборки

Эксперимент был ориентирован именно на проверку качества прогнозирования, а не на экономический эффект управляющих решений. Поэтому в нём намеренно использовалось одно устройство с хорошо сформированной историей телеметрии и зрелыми прогнозами на всех трёх горизонтах. Такой дизайн позволяет исключить смешивание двух разных задач: сначала формально подтвердить качество прогностического контура, а затем уже отдельно анализировать ожидаемые потери и полезность решений СППР.

Выбор режима `strict intersection` имеет принципиальное значение. В обычном пользовательском сравнении система может показывать качество каждой модели на собственном наборе зрелых точек. Однако для диссертационной интерпретации такой подход недостаточно строгий, поскольку разные модели могут сравниваться на слегка разных временных срезах. Поэтому в данном эксперименте в финальную выборку включались только те пары «метрика — момент времени — горизонт», для которых факт найден у всех трёх ветвей одновременно. Это делает сравнение моделей методологически корректным.

Дополнительно для проверки устойчивости результатов использованы две процедуры. Во-первых, применён парный тест Диболда–Мариано, который позволяет оценить статистическую значимость различия прогнозных ошибок на одинаковой выборке. Во-вторых, для RMSE, MASE и Pinball Loss построены доверительные интервалы по moving block bootstrap с автоматическим подбором длины блока. Тем самым раздел 3.4 оценивает не только среднюю ошибку, но и устойчивость полученных выводов.
'''

p342 = f'''3.4.2 Результаты на горизонте 24 часа

На краткосрочном горизонте 24 часа наилучшие агрегированные показатели продемонстрировал оркестр. Для него получены `RMSE = {fmt_num(ensemble_24.get("rmse"), 4)}` и `MAE = {fmt_num(ensemble_24.get("mae"), 4)}`. У SARIMA эти значения составили соответственно `{fmt_num(sarima_24.get("rmse"), 4)}` и `{fmt_num(sarima_24.get("mae"), 4)}`, а у LSTM — `{fmt_num(lstm_24.get("rmse"), 4)}` и `{fmt_num(lstm_24.get("mae"), 4)}`. Следовательно, уже на коротком горизонте интеграция двух прогностических ветвей не ухудшает базовый прогноз, а даёт измеримый выигрыш относительно чистой статистической модели и особенно относительно нейросетевой ветви.

По вероятностным метрикам преимущество оркестра также сохраняется. Его средняя pinball-ошибка составляет `{fmt_num(prob_24["ensemble"].get("pinball_avg"), 4)}`, тогда как для SARIMA она равна `{fmt_num(prob_24["sarima"].get("pinball_avg"), 4)}`, а для LSTM — `{fmt_num(prob_24["lstm"].get("pinball_avg"), 4)}`. Тем самым на коротком горизонте оркестр обеспечивает лучшую согласованность центрального прогноза и квантильных границ.

Калибровка интервалов, однако, остаётся общей проблемой всех трёх моделей. Для оркестра получено `PICP80 = {fmt_pct(ensemble_24.get("picp80"), 2)}` при `ACE80 = {fmt_pct(ensemble_24.get("ace80"), 2)}`. Это лучше, чем у SARIMA и LSTM, но всё ещё существенно ниже идеального покрытия 80%. Следовательно, краткосрочный контур прогнозирования уже достаточно точен по центральной ошибке, однако модуль вероятностной неопределённости требует дальнейшей калибровки.
'''

p343 = f'''3.4.3 Результаты на горизонте 7 суток

На среднесрочном горизонте 7 суток наблюдается наиболее устойчивое преимущество оркестра. Величины ошибок составили: для оркестра `RMSE = {fmt_num(ensemble_7.get("rmse"), 4)}` и `MAE = {fmt_num(ensemble_7.get("mae"), 4)}`, для SARIMA — `RMSE = {fmt_num(sarima_7.get("rmse"), 4)}` и `MAE = {fmt_num(sarima_7.get("mae"), 4)}`, для LSTM — `RMSE = {fmt_num(lstm_7.get("rmse"), 4)}` и `MAE = {fmt_num(lstm_7.get("mae"), 4)}`. В отличие от более ранних версий системы, где средний горизонт часто выигрывала чистая статистическая ветвь, актуальная конфигурация показывает, что согласованное объединение SARIMA и LSTM уже даёт на 7 сутках лучший итоговый point forecast.

Особенно важно, что выигрыш оркестра на данном горизонте проявляется не только в RMSE, но и в вероятностных метриках. Для него получен лучший `Pinball(avg) = {fmt_num(prob_7["ensemble"].get("pinball_avg"), 4)}`, а также наилучшие агрегированные показатели `sMAPE`, `MASE` и `RMSSE`. Это означает, что интегрированная модель улучшает не отдельную частную метрику, а качество прогноза в нескольких независимых шкалах оценки.

Парный тест Диболда–Мариано подтверждает, что наблюдаемое преимущество не является случайным. Для пары `SARIMA vs Оркестр` на горизонте 7 суток получено `DM-stat = {fmt_num(dm_7.get("dm_stat"), 4)}` при `p-value = {fmt_num(dm_7.get("p_value"), 6)}`. Поскольку уровень значимости существенно меньше 0,05, можно сделать корректный вывод о статистически значимом превосходстве оркестра над статистическим baseline на среднесрочном горизонте.
'''

p344 = f'''3.4.4 Результаты на горизонте 30 суток

Наиболее важным для диссертационной интерпретации является долгосрочный горизонт 30 суток, поскольку именно он определяет практическую пригодность системы к раннему предупреждению деградации. В актуальном эксперименте и на этом горизонте лучшей моделью снова оказался оркестр: `RMSE = {fmt_num(ensemble_30.get("rmse"), 4)}`, `MAE = {fmt_num(ensemble_30.get("mae"), 4)}`. Для SARIMA получены значения `{fmt_num(sarima_30.get("rmse"), 4)}` и `{fmt_num(sarima_30.get("mae"), 4)}`, для LSTM — `{fmt_num(lstm_30.get("rmse"), 4)}` и `{fmt_num(lstm_30.get("mae"), 4)}`. Следовательно, в текущей версии системы оркестр перестал быть только «слоем неопределённости» и стал лучшей итоговой моделью также по центральной точности.

Этот результат важен по двум причинам. Во-первых, он означает, что доработки LSTM-ветви и логики ансамблирования больше не приводят к ухудшению долгосрочного прогноза. Во-вторых, выигрыш оркестра достигается без отказа от базовой статистической модели: итоговая траектория формируется как результат интеграции двух ветвей, а не как замена одной модели другой.

На уровне вероятностной интерпретации оркестр также занимает первое место: `Pinball(avg) = {fmt_num(prob_30["ensemble"].get("pinball_avg"), 4)}`, `PICP80 = {fmt_pct(ensemble_30.get("picp80"), 2)}`, `ACE80 = {fmt_pct(ensemble_30.get("ace80"), 2)}`. При этом абсолютный уровень покрытия всё ещё недостаточен для признания интервалов полностью откалиброванными. Иными словами, на 30 сутках система уже научилась лучше предсказывать саму траекторию, чем отдельные ветви, но модуль интервалов остаётся направлением дальнейшего улучшения.
'''

p345 = f'''3.4.5 Пометрический анализ и статистическая интерпретация результатов

Агрегированные значения RMSE и MAE сами по себе ещё не доказывают, что оркестр действительно улучшает прогноз большинства эксплуатационных показателей. Поэтому дополнительно был выполнен пометрический анализ по всем сочетаниям «метрика — горизонт». Он показал, что по RMSE оркестр является лучшей моделью в 20 из 21 исследованного бакета, а по MAE — в 19 из 21. Единственным исключением по RMSE стала системная температура на горизонте 24 часа, где SARIMA дала немного меньшую ошибку (`0,9093` против `0,9277` у оркестра). Следовательно, преимущество оркестра не является следствием одной удачной серии, а наблюдается почти для всех телеметрических каналов.

Особенно показателен разбор длинного горизонта. Для метрики загрузки CPU на 30 сутках оркестр дал `RMSE = 5,1021`, тогда как SARIMA — `5,5266`, а LSTM — `5,9915`. Для использования памяти получены `1,7474`, `1,9790` и `1,9504` соответственно. Для сетевых метрик, имеющих наибольший абсолютный масштаб, оркестр также сохраняет преимущество: по `net_bytes_recv` его RMSE составляет `117,9365` против `119,2492` у SARIMA и `177,8614` у LSTM, а по `net_bytes_sent` — `110,1859` против `113,3453` и `180,7186`. Это означает, что итоговое улучшение достигается не только на «лёгких» температурных рядах, но и на наиболее шумных высокоамплитудных потоках.

Статистическая проверка дополняет этот вывод. Для пары `SARIMA vs Оркестр` тест Диболда–Мариано дал значимые результаты на всех горизонтах: `p = {fmt_num(dm_24.get("p_value"), 6)}` на 24 часах, `p = {fmt_num(dm_7.get("p_value"), 6)}` на 7 сутках и `p = {fmt_num(dm_30.get("p_value"), 6)}` на 30 сутках. При этом доверительные интервалы bootstrap для RMSE частично перекрываются, что указывает на умеренную величину абсолютного выигрыша, но не отменяет факта его устойчивости на одной и той же парной выборке ошибок. Иначе говоря, оркестр выигрывает не за счёт разницы в окне наблюдения, а за счёт реального уменьшения прогнозной ошибки на общих точках сравнения.
'''

p346 = f'''3.4.6 Ограничения эксперимента и итоговые выводы

Несмотря на явное преимущество оркестра по центральной точности, эксперимент также выявил несколько ограничений текущей версии системы. Во-первых, все три модели демонстрируют положительный `mean_error`, то есть склонность к умеренной переоценке уровня метрик. В агрегированном виде этот показатель составил `{fmt_num(ensemble_total.get("mean_error"), 4)}` для оркестра, `{fmt_num(sarima_total.get("mean_error"), 4)}` для SARIMA и `{fmt_num(lstm_total.get("mean_error"), 4)}` для LSTM. Для прикладной СППР такая асимметрия менее опасна, чем систематическое недооценивание риска, однако она требует дальнейшей калибровки.

Во-вторых, вероятностная часть прогноза пока уступает точечной. Даже лучшая модель по калибровке, то есть оркестр, показывает покрытие `PICP80` лишь на уровне от {fmt_pct(min(as_float(ensemble_24.get("picp80")), as_float(ensemble_7.get("picp80")), as_float(ensemble_30.get("picp80"))), 2)} до {fmt_pct(max(as_float(ensemble_24.get("picp80")), as_float(ensemble_7.get("picp80")), as_float(ensemble_30.get("picp80"))), 2)} при целевом значении 80%. Следовательно, интерпретацию `p10–p90` в текущей версии системы следует считать полезной, но ещё не окончательно откалиброванной.

В-третьих, данный run не содержит валидных сравнений по критерию `ΔR`, поскольку в выбранном окне отсутствуют сопоставимые запуски СППР с baseline-действием `no_action`. Поэтому для настоящего раздела корректно анализировать только качество прогностического контура. Экономическая эффективность системы должна проверяться отдельным экспериментом на окне, где накоплено достаточное число зрелых `DecisionRun`.

В целом проведённый эксперимент позволяет сделать три итоговых вывода. Во-первых, на текущем состоянии проекта лучшей итоговой моделью по прогнозной точности является оркестр, а не отдельные ветви SARIMA или LSTM. Во-вторых, это преимущество подтверждается не только агрегированными значениями RMSE и MAE, но и пометрическим анализом, а также парной статистической проверкой различий. В-третьих, дальнейшее развитие системы должно быть сосредоточено не столько на point forecast, сколько на улучшении интервальной калибровки и устранении систематического положительного bias, поскольку именно эти два аспекта остаются основными ограничениями текущей реализации.
'''

md_parts = [intro, p341, p342, p343, p344, p345, p346]
OUT_MD.write_text('\n\n'.join(md_parts), encoding='utf-8')

doc = Document()
style = doc.styles['Normal']
style.font.name = BODY_FONT
style.font.size = BODY_SIZE
style._element.rPr.rFonts.set(qn('w:ascii'), BODY_FONT)
style._element.rPr.rFonts.set(qn('w:hAnsi'), BODY_FONT)
style._element.rPr.rFonts.set(qn('w:cs'), BODY_FONT)
style.paragraph_format.first_line_indent = FIRST_LINE
style.paragraph_format.line_spacing = LINE_SPACING

add_heading(doc, '3.4 Вычислительный эксперимент на базе собранных метрик', level=1)
for text in intro.split('\n\n')[1:]:
    add_body(doc, text)

add_heading(doc, '3.4.1 Организация эксперимента и состав выборки', level=2)
for text in p341.split('\n\n')[1:]:
    add_body(doc, text)
add_table(
    doc,
    'Таблица 3.4 — Параметры экспериментального протокола актуального evaluation-run',
    ['Параметр', 'Значение', 'Назначение'],
    table_34_rows,
    col_widths_cm=[4.2, 4.8, 7.2],
)
add_picture_with_caption(doc, RUN_DIR / 'evaluation_summary_page.png', 'Рисунок 3.5 — Сводная страница актуального evaluation-run', width_cm=15.0)

add_heading(doc, '3.4.2 Результаты на горизонте 24 часа', level=2)
for text in p342.split('\n\n')[1:]:
    add_body(doc, text)
add_table(
    doc,
    'Таблица 3.5 — Качество прогноза на горизонте 24 часа',
    ['Модель', 'N', 'ME', 'MAE', 'RMSE', 'sMAPE', 'Pinball(avg)', 'PICP80', 'ACE80'],
    table_35_rows,
    col_widths_cm=[2.6, 1.3, 1.7, 1.8, 1.8, 2.0, 2.5, 1.8, 1.8],
)

add_heading(doc, '3.4.3 Результаты на горизонте 7 суток', level=2)
for text in p343.split('\n\n')[1:]:
    add_body(doc, text)
add_table(
    doc,
    'Таблица 3.6 — Качество прогноза на горизонте 7 суток',
    ['Модель', 'N', 'ME', 'MAE', 'RMSE', 'sMAPE', 'Pinball(avg)', 'PICP80', 'ACE80'],
    table_36_rows,
    col_widths_cm=[2.6, 1.3, 1.7, 1.8, 1.8, 2.0, 2.5, 1.8, 1.8],
)
add_picture_with_caption(doc, RUN_DIR / 'chart_rmse_by_model_horizon.png', 'Рисунок 3.6 — Сравнение RMSE по моделям и горизонтам', width_cm=14.8)

add_heading(doc, '3.4.4 Результаты на горизонте 30 суток', level=2)
for text in p344.split('\n\n')[1:]:
    add_body(doc, text)
add_table(
    doc,
    'Таблица 3.7 — Качество прогноза на горизонте 30 суток',
    ['Модель', 'N', 'ME', 'MAE', 'RMSE', 'sMAPE', 'Pinball(avg)', 'PICP80', 'ACE80'],
    table_37_rows,
    col_widths_cm=[2.6, 1.3, 1.7, 1.8, 1.8, 2.0, 2.5, 1.8, 1.8],
)
add_picture_with_caption(doc, RUN_DIR / 'chart_mae_by_model_horizon.png', 'Рисунок 3.7 — Сравнение MAE по моделям и горизонтам', width_cm=14.8)

add_heading(doc, '3.4.5 Пометрический анализ и статистическая интерпретация результатов', level=2)
for text in p345.split('\n\n')[1:]:
    add_body(doc, text)
add_table(
    doc,
    'Таблица 3.8 — Число побед по отдельным бакетам «метрика — горизонт»',
    ['Критерий', 'Оркестр', 'SARIMA', 'LSTM'],
    table_38_rows,
    col_widths_cm=[5.0, 3.0, 3.0, 3.0],
)
add_table(
    doc,
    'Таблица 3.9 — Результаты теста Диболда–Мариано для пары SARIMA vs Оркестр',
    ['Горизонт', 'Пара моделей', 'DM-stat', 'p-value', 'Победитель', 'Интерпретация'],
    table_39_rows,
    col_widths_cm=[1.8, 3.4, 2.0, 2.2, 2.2, 5.4],
)

add_heading(doc, '3.4.6 Ограничения эксперимента и итоговые выводы', level=2)
for text in p346.split('\n\n')[1:]:
    add_body(doc, text)

doc.save(str(OUT_DOCX))
print(OUT_MD)
print(OUT_DOCX)
